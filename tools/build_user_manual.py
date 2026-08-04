from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


FONT_EAST_ASIA = "Microsoft YaHei"
FONT_LATIN = "Arial"
ACCENT = "0E7490"
ACCENT_DARK = "155E75"
TEXT = "1F2937"
MUTED = "64748B"
PLACEHOLDER_FILL = "E8F3F7"
WARNING_FILL = "FFF4E5"
QUOTE_FILL = "F3F4F6"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=120, start=140, bottom=120, end=140) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def keep_row_together(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_run_font(run, size: float | None = None, bold: bool | None = None, color: str | None = None) -> None:
    run.font.name = FONT_LATIN
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT_EAST_ASIA)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


INLINE_RE = re.compile(r"(\*\*.+?\*\*|`.+?`)")


def add_inline(paragraph, text: str, *, base_size: float = 10.5, base_color: str = TEXT) -> None:
    for part in INLINE_RE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            set_run_font(run, base_size, bold=True, color=base_color)
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            set_run_font(run, base_size - 0.5, color=ACCENT_DARK)
            run.font.name = "Consolas"
            run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        else:
            run = paragraph.add_run(part)
            set_run_font(run, base_size, color=base_color)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("第 ")
    set_run_font(run, 9, color=MUTED)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instr, end))
    tail = paragraph.add_run(" 页")
    set_run_font(tail, 9, color=MUTED)


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Cm(2.1)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.0)
    section.header_distance = Cm(0.8)
    section.footer_distance = Cm(0.8)

    normal = document.styles["Normal"]
    normal.font.name = FONT_LATIN
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_EAST_ASIA)
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(TEXT)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.2

    heading_sizes = {1: 22, 2: 16, 3: 12.5}
    for level, size in heading_sizes.items():
        style = document.styles[f"Heading {level}"]
        style.font.name = FONT_LATIN
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_EAST_ASIA)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(ACCENT_DARK if level > 1 else ACCENT)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(12 if level > 1 else 0)
        style.paragraph_format.space_after = Pt(6)

    for style_name in ("List Bullet", "List Number"):
        style = document.styles[style_name]
        style.font.name = FONT_LATIN
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_EAST_ASIA)
        style.font.size = Pt(10.5)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header.add_run("NewLight_Analysis 用户使用手册")
    set_run_font(run, 9, color=MUTED)
    add_page_number(section.footer.paragraphs[0])

    settings = document.settings._element
    update_fields = OxmlElement("w:updateFields")
    update_fields.set(qn("w:val"), "true")
    settings.append(update_fields)


def add_callout(document: Document, text: str) -> None:
    is_placeholder = text.startswith("[图片占位")
    is_warning = "DeepCAD-RT" in text or "硬件警告" in text
    table = document.add_table(rows=1, cols=1)
    table.autofit = True
    cell = table.cell(0, 0)
    set_cell_shading(cell, PLACEHOLDER_FILL if is_placeholder else WARNING_FILL if is_warning else QUOTE_FILL)
    set_cell_margins(cell)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    row = table.rows[0]
    keep_row_together(row)
    if is_placeholder:
        row.height = Inches(0.8)
        row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if is_placeholder else WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_after = Pt(0)
    add_inline(paragraph, text, base_size=10, base_color=ACCENT_DARK if is_placeholder else TEXT)
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def parse_table(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    rows: list[list[str]] = []
    index = start
    while index < len(lines) and lines[index].strip().startswith("|"):
        row = [cell.strip() for cell in lines[index].strip().strip("|").split("|")]
        rows.append(row)
        index += 1
    if len(rows) >= 2 and all(re.fullmatch(r":?-{3,}:?", cell.replace(" ", "")) for cell in rows[1]):
        del rows[1]
    return rows, index


def add_table(document: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    columns = max(len(row) for row in rows)
    table = document.add_table(rows=len(rows), cols=columns)
    table.style = "Table Grid"
    table.autofit = True
    for row_index, source_row in enumerate(rows):
        for column_index in range(columns):
            cell = table.cell(row_index, column_index)
            text = source_row[column_index] if column_index < len(source_row) else ""
            set_cell_margins(cell, top=75, start=90, bottom=75, end=90)
            if row_index == 0:
                set_cell_shading(cell, ACCENT_DARK)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            add_inline(paragraph, text, base_size=8.5, base_color="FFFFFF" if row_index == 0 else TEXT)
            if row_index == 0:
                for run in paragraph.runs:
                    run.bold = True
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def add_code_block(document: Document, code_lines: list[str]) -> None:
    table = document.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    set_cell_shading(cell, "EFF6FF")
    set_cell_margins(cell, top=100, start=120, bottom=100, end=120)
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    for index, line in enumerate(code_lines):
        if index:
            paragraph.add_run("\n")
        run = paragraph.add_run(line)
        set_run_font(run, 8.5, color="1E3A5F")
        run.font.name = "Consolas"
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def build_docx(source: Path, output: Path) -> None:
    lines = source.read_text(encoding="utf-8").splitlines()
    document = Document()
    configure_document(document)

    first_rule = True
    index = 0
    paragraph_buffer: list[str] = []

    def flush_paragraph() -> None:
        if not paragraph_buffer:
            return
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.first_line_indent = Cm(0.74)
        add_inline(paragraph, " ".join(item.strip() for item in paragraph_buffer))
        paragraph_buffer.clear()

    while index < len(lines):
        raw = lines[index]
        stripped = raw.strip()
        if stripped.startswith("```"):
            flush_paragraph()
            index += 1
            code: list[str] = []
            while index < len(lines) and not lines[index].strip().startswith("```"):
                code.append(lines[index])
                index += 1
            add_code_block(document, code)
            index += 1
            continue
        if stripped.startswith("|"):
            flush_paragraph()
            rows, index = parse_table(lines, index)
            add_table(document, rows)
            continue
        if stripped.startswith("> "):
            flush_paragraph()
            add_callout(document, stripped[2:].strip())
            index += 1
            continue
        heading = re.match(r"^(#{1,3})\s+(.+)$", stripped)
        if heading:
            flush_paragraph()
            level = len(heading.group(1))
            text = heading.group(2)
            paragraph = document.add_heading(level=level)
            add_inline(paragraph, text, base_size={1: 22, 2: 16, 3: 12.5}[level], base_color=ACCENT if level == 1 else ACCENT_DARK)
            if level == 1:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.paragraph_format.space_before = Pt(80)
                paragraph.paragraph_format.space_after = Pt(18)
            index += 1
            continue
        if stripped == "---":
            flush_paragraph()
            if first_rule:
                document.add_page_break()
                first_rule = False
            index += 1
            continue
        bullet = re.match(r"^-\s+(.+)$", stripped)
        numbered = re.match(r"^\d+\.\s+(.+)$", stripped)
        if bullet or numbered:
            flush_paragraph()
            paragraph = document.add_paragraph(style="List Bullet" if bullet else "List Number")
            add_inline(paragraph, (bullet or numbered).group(1))
            index += 1
            continue
        if not stripped:
            flush_paragraph()
            index += 1
            continue
        paragraph_buffer.append(stripped.removesuffix("  "))
        index += 1

    flush_paragraph()
    output.parent.mkdir(parents=True, exist_ok=True)
    document.core_properties.title = "NewLight_Analysis 用户使用手册"
    document.core_properties.subject = "双光子与钙成像分析软件操作、参数、输出和方法学说明"
    document.core_properties.author = "NewLight_Analysis Project"
    document.save(output)


def main() -> int:
    parser = argparse.ArgumentParser(description="Build the NewLight_Analysis DOCX manual from Markdown.")
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    build_docx(args.source.resolve(), args.output.resolve())
    print(args.output.resolve())
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
